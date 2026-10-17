import asyncio
import shutil
import sys
import threading
import time
from collections.abc import Callable, Iterator
from pathlib import Path

import pytest

# Make `app` importable when running pytest from backend/
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.config import Settings  # noqa: E402
from app.tts import audio  # noqa: E402

WORDS = "zero one two three four five six seven eight nine ten eleven twelve".split()


async def no_sleep(_: float) -> None:
    return None


class GateEngine:
    """Fake TTS engine for API / job tests: copies a valid MP3 clip. `gate` (if given) blocks every
    call until it is set; `fails(text)` makes a call raise; `calls` records the texts."""

    name = "edge"

    def __init__(
        self,
        clip: Path,
        fails: Callable[[str], bool] = lambda _: False,
        gate: threading.Event | None = None,
        delay: float = 0.0,
    ) -> None:
        self.clip = clip
        self.fails = fails
        self.gate = gate
        self.delay = delay
        self.calls: list[str] = []

    async def synthesize(self, text: str, out_path: Path, voice: str | None = None) -> None:
        self.calls.append(text)
        if self.gate is not None:
            await asyncio.to_thread(self.gate.wait, 30)
        if self.delay:
            await asyncio.sleep(self.delay)
        if self.fails(text):
            raise ConnectionError("fake edge failure")
        shutil.copyfile(self.clip, out_path)


@pytest.fixture(scope="session")
def mp3_clip(tmp_path_factory: pytest.TempPathFactory) -> Path:
    if shutil.which("ffmpeg") is None:
        pytest.skip("ffmpeg not on PATH")
    return audio.make_silence(240, tmp_path_factory.mktemp("clip") / "clip.mp3")


@pytest.fixture
def api_settings(tmp_path: Path) -> Settings:
    return Settings(
        _env_file=None,
        storage_dir=tmp_path / "storage",
        max_file_mb=1,
        max_pdf_pages=5,
        tts_concurrency=2,
        tts_fallback="none",
        tts_max_retries=1,
        cleanup_interval_minutes=60,
    )


def make_docx(path: Path, chapters: int = 3, words_per_chapter: int = 60, title: str = "Chapter") -> Path:
    """A DOCX with Heading 1 chapters (structure from styles); numbers spelled out."""
    import docx

    document = docx.Document()
    for n in range(1, chapters + 1):
        document.add_heading(f"{title} {WORDS[n].capitalize()}", level=1)
        sentence = f"This is chapter {WORDS[n]} and it has some words to read aloud."
        repeats = max(1, words_per_chapter // len(sentence.split()))
        for _ in range(2):
            document.add_paragraph(" ".join([sentence] * (repeats // 2 or 1)))
    document.save(path)
    return path


def wait_for(predicate: Callable[[], bool], timeout: float = 30.0, step: float = 0.02) -> None:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if predicate():
            return
        time.sleep(step)
    raise AssertionError("condition not met in time")


@pytest.fixture
def gate() -> Iterator[threading.Event]:
    event = threading.Event()
    yield event
    event.set()  # never leave a worker blocked after a test
