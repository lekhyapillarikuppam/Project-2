"""Character-voices mode through the API (SPEC §15.1) with a fake TTS engine and a fake Groq
(`httpx.MockTransport`): mode validation, the job object, attribution one chapter ahead, the
baseline fallback when the LLM is down, and single mode unchanged."""

import json
import threading
from collections.abc import Callable, Iterator
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient

from app.config import Settings
from app.dialogue.llm import GroqClient, TokenLimiter
from app.jobs.registry import JobRegistry
from app.main import create_app
from conftest import WORDS, GateEngine, wait_for

FINISHED = ("done", "partial", "failed")


def dialogue_docx(path: Path, chapters: int = 3) -> Path:
    """Heading 1 chapters of dialogue between Alice and the Hatter (curly single quotes)."""
    import docx

    document = docx.Document()
    for n in range(1, chapters + 1):
        word = WORDS[n]
        document.add_heading(f"Chapter {word.capitalize()}", level=1)
        document.add_paragraph(f"The tea party went on in chapter {word}, and nobody seemed to mind.")
        document.add_paragraph(f"‘Have some wine,’ said the Hatter. ‘It is chapter {word} after all.’")
        document.add_paragraph("‘I don’t see any wine,’ Alice remarked.")
        document.add_paragraph("‘There isn’t any,’ said the Hatter.")
        document.add_paragraph("‘Then it wasn’t very civil of you to offer it,’ said Alice angrily.")
    document.save(path)
    return path


class FakeGroq:
    """Answers every quote id with a speaker; records which chapter each request was for."""

    def __init__(self, status: int = 200) -> None:
        self.status = status
        self.chapters: list[str] = []
        self.lock = threading.Lock()

    def __call__(self, request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        user = body["messages"][1]["content"]
        with self.lock:
            self.chapters.append(next(w for w in WORDS if f"chapter {w}" in user.split("TEXT:")[-1]))
        if self.status != 200:
            return httpx.Response(self.status, text="down")
        ids = sorted({part.split("]")[0] for part in user.split("[Q")[1:]}, key=int)
        q = {f"Q{i}": ("Hatter" if int(i) % 2 else "Alice") for i in ids}
        new = [{"name": "Alice", "gender": "female", "aliases": []}, {"name": "Hatter", "gender": "male", "aliases": []}]
        content = json.dumps({"q": q, "new": new})
        return httpx.Response(200, json={"choices": [{"message": {"content": content}}],
                                         "usage": {"prompt_tokens": 100, "completion_tokens": 20}})  # fmt: skip


def llm_factory(groq: FakeGroq) -> Callable[[Settings], GroqClient]:
    def build(settings: Settings) -> GroqClient:
        limiter = TokenLimiter(10**6, 10**6)
        return GroqClient("test-key", settings.llm_model, limiter, transport=httpx.MockTransport(groq))

    return build


@pytest.fixture
def char_settings(api_settings: Settings) -> Settings:
    return api_settings.model_copy(update={"groq_api_key": "test-key", "attribution_method": "llm"})


def make_client(settings: Settings, engine: GateEngine, groq: FakeGroq) -> TestClient:
    return TestClient(create_app(settings, lambda s: engine, llm_factory(groq)))


@pytest.fixture
def groq() -> FakeGroq:
    return FakeGroq()


@pytest.fixture
def client(char_settings: Settings, mp3_clip: Path, groq: FakeGroq, gate: threading.Event) -> Iterator[TestClient]:
    with make_client(char_settings, GateEngine(mp3_clip, gate=gate), groq) as c:
        yield c
    gate.set()


def upload(client: TestClient, path: Path, mode: str | None = None) -> httpx.Response:
    data = {"narration_mode": mode} if mode is not None else {}
    with path.open("rb") as fh:
        return client.post("/api/upload", files={"file": (path.name, fh, "application/octet-stream")}, data=data)


def wait_done(client: TestClient, job_id: str) -> dict:
    result: dict = {}

    def finished() -> bool:
        result.update(client.get(f"/api/jobs/{job_id}").json())
        return result.get("status") in FINISHED

    wait_for(finished)
    return result


def test_mode_validation(client: TestClient, char_settings: Settings, mp3_clip: Path, tmp_path: Path) -> None:
    book = dialogue_docx(tmp_path / "book.docx")
    r = upload(client, book, "chorus")
    assert r.status_code == 400 and r.json()["error"]["code"] == "BAD_REQUEST"
    assert "narration_mode" in r.json()["error"]["message"]
    jobs = char_settings.storage_path / "jobs"
    assert not any(jobs.iterdir())  # a rejected upload leaves nothing behind

    no_key = char_settings.model_copy(update={"groq_api_key": ""})
    with make_client(no_key, GateEngine(mp3_clip), FakeGroq()) as c:
        r = upload(c, book, "characters")
        assert r.status_code == 400 and "Character voices aren't available" in r.json()["error"]["message"]
        assert c.get("/api/config").json()["character_voices_available"] is False
    assert client.get("/api/config").json()["character_voices_available"] is True


def test_characters_job_object(
    client: TestClient, tmp_path: Path, gate: threading.Event, monkeypatch: pytest.MonkeyPatch
) -> None:
    history: list[dict] = []
    original = JobRegistry._persist

    def spy(job):  # noqa: ANN001, ANN202
        history.append(job.to_public())
        original(job)

    monkeypatch.setattr(JobRegistry, "_persist", staticmethod(spy))
    r = upload(client, dialogue_docx(tmp_path / "book.docx"), "characters")
    assert r.status_code == 202
    gate.set()
    final = wait_done(client, r.json()["job_id"])

    statuses = [h["status"] for h in history]
    deduped = [s for i, s in enumerate(statuses) if i == 0 or statuses[i - 1] != s]
    assert deduped == ["queued", "extracting", "attributing", "synthesizing", "done"]
    assert final["narration_mode"] == "characters"
    progress = final["progress"]
    assert progress["windows_total"] == progress["windows_done"] == 3
    assert [c["attribution"] for c in final["chapters"]] == ["llm"] * 3
    names = {c["name"]: c for c in final["cast"]}
    assert set(names) == {"Alice", "Hatter"}
    assert names["Alice"]["gender"] == "female" and names["Hatter"]["gender"] == "male"
    assert names["Alice"]["voice"] != names["Hatter"]["voice"]
    assert names["Alice"]["voice_label"].startswith("Voice ")
    assert not any(w.startswith("Speaker detection") for w in final["warnings"])


def test_next_chapter_is_attributed_while_the_current_one_is_narrated(
    client: TestClient, groq: FakeGroq, tmp_path: Path, gate: threading.Event
) -> None:
    job_id = upload(client, dialogue_docx(tmp_path / "book.docx"), "characters").json()["job_id"]
    # The engine is blocked inside chapter 1; chapter 2's attribution must still happen.
    wait_for(lambda: client.get(f"/api/jobs/{job_id}").json()["progress"]["windows_done"] == 2)
    job = client.get(f"/api/jobs/{job_id}").json()
    assert job["chapters"][0]["status"] == "generating" and job["chapters"][1]["status"] == "pending"
    assert groq.chapters == ["one", "two"]  # not three: only one chapter ahead
    gate.set()
    assert wait_done(client, job_id)["status"] == "done"
    assert groq.chapters == ["one", "two", "three"]


def test_llm_down_gives_baseline_chapters_and_a_warning(
    char_settings: Settings, mp3_clip: Path, tmp_path: Path
) -> None:
    with make_client(char_settings, GateEngine(mp3_clip), FakeGroq(status=500)) as c:
        job_id = upload(c, dialogue_docx(tmp_path / "book.docx"), "characters").json()["job_id"]
        final = wait_done(c, job_id)
    assert final["status"] == "done"
    assert [ch["attribution"] for ch in final["chapters"]] == ["baseline"] * 3
    assert final["progress"]["windows_done"] == 3
    assert any("rule-based method for 3 chapters" in w for w in final["warnings"])
    assert {m["name"] for m in final["cast"]} == {"Alice", "Hatter"}  # explicit tags still name them


def test_single_mode_is_unchanged(client: TestClient, groq: FakeGroq, tmp_path: Path, gate: threading.Event) -> None:
    job_id = upload(client, dialogue_docx(tmp_path / "book.docx")).json()["job_id"]
    gate.set()
    final = wait_done(client, job_id)
    assert final["status"] == "done" and final["narration_mode"] == "single"
    assert final["cast"] == [] and final["progress"]["windows_total"] == 0
    assert [c["attribution"] for c in final["chapters"]] == ["none"] * 3
    assert groq.chapters == []  # the LLM is never called
